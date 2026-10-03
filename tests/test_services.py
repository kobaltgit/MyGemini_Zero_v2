"""
Unit tests for services:
- services/error_parser.py (Gemini error key mapping)
- services/media.py (Multimodal image/audio Part packaging & document parsing)
- services/dialog_namer.py (Fallback title extraction)
- services/calendar_helper.py (Interactive calendar inline keyboard generation)
"""

import io
import datetime
import pytest
from docx import Document
from google.genai import types
from services.error_parser import get_user_friendly_error_key, ERROR_MAP, DEFAULT_ERROR_KEY
from services.media import format_image_part, format_audio_part, parse_document_text
from services.dialog_namer import fallback_title_from_text
from services.calendar_helper import (
    create_calendar_keyboard,
    CALLBACK_CALENDAR_DATE_PREFIX,
    CALLBACK_CALENDAR_MONTH_PREFIX,
)
from services.gemini import model_supports_search


class TestModelSupportsSearch:
    """Tests for model_supports_search filtering."""

    def test_flash_lite_search_disabled(self):
        assert model_supports_search("gemini-2.5-flash-lite") is False
        assert model_supports_search("gemini-2.0-flash-lite") is False
        assert model_supports_search("models/gemini-2.5-flash-lite-preview-06-17") is False

    def test_standard_models_search_enabled(self):
        assert model_supports_search("gemini-2.5-flash") is True
        assert model_supports_search("gemini-2.5-pro") is True
        assert model_supports_search("gemini-2.0-flash") is True

    def test_non_gemini_or_specialized_disabled(self):
        assert model_supports_search("gemini-embedding-001") is False
        assert model_supports_search("imagen-3.0") is False



class TestErrorParser:
    """Tests for services/error_parser.py."""

    def test_quota_exceeded_429(self):
        err = {"code": 429, "message": "Resource has been exhausted (e.g. check quota)."}
        key = get_user_friendly_error_key(err)
        assert key == "gemini_error_quota_exceeded"

    def test_api_key_invalid(self):
        err = {"error": {"code": 400, "status": "API_KEY_INVALID", "message": "API key not valid."}}
        key = get_user_friendly_error_key(err)
        assert key == "gemini_error_api_key_invalid"

    def test_safety_block(self):
        err = {"candidates": [{"finish_reason": "SAFETY"}]}
        key = get_user_friendly_error_key(err)
        assert key == "gemini_error_safety"

    def test_service_unavailable_503(self):
        err = {"error": "The model is overloaded. Please try again later. service_unavailable"}
        key = get_user_friendly_error_key(err)
        assert key == "gemini_error_unavailable"

    def test_search_tool_unsupported(self):
        err = {"message": "tool is not supported for this model"}
        key = get_user_friendly_error_key(err)
        assert key == "gemini_error_search_not_supported"

    def test_unknown_error_fallback(self):
        err = {"random_key": "some completely unknown error 999"}
        key = get_user_friendly_error_key(err)
        assert key == DEFAULT_ERROR_KEY

        # Non-dict inputs
        assert get_user_friendly_error_key(None) == DEFAULT_ERROR_KEY
        assert get_user_friendly_error_key("error string") == DEFAULT_ERROR_KEY


class TestMediaService:
    """Tests for services/media.py."""

    def test_format_image_part(self):
        img_bytes = b"\xFF\xD8\xFF\xE0FakeJpegBytes"
        part = format_image_part(img_bytes, mime_type="image/jpeg")
        assert isinstance(part, types.Part)
        assert part.inline_data.data == img_bytes
        assert part.inline_data.mime_type == "image/jpeg"

    def test_format_audio_part(self):
        audio_bytes = b"OggS\x00FakeAudioBytes"
        part = format_audio_part(audio_bytes, mime_type="audio/ogg")
        assert isinstance(part, types.Part)
        assert part.inline_data.data == audio_bytes
        assert part.inline_data.mime_type == "audio/ogg"

    def test_parse_document_text_plain(self):
        text_content = "Hello, world! Это тестовый файл.\nСтрока 2."
        file_bytes = text_content.encode("utf-8")
        extracted, doc_type = parse_document_text(file_bytes, "notes.txt")
        assert doc_type == "text"
        assert extracted == text_content

    def test_parse_document_text_docx(self):
        # Create an in-memory DOCX
        doc = Document()
        doc.add_paragraph("First paragraph in DOCX.")
        doc.add_paragraph("Second paragraph in DOCX.")
        buf = io.BytesIO()
        doc.save(buf)
        file_bytes = buf.getvalue()

        extracted, doc_type = parse_document_text(file_bytes, "report.docx")
        assert doc_type == "docx"
        assert "First paragraph in DOCX." in extracted
        assert "Second paragraph in DOCX." in extracted

    def test_parse_document_empty(self):
        extracted, doc_type = parse_document_text(b"", "empty.txt")
        assert doc_type == "text"
        assert extracted == ""


class TestDialogNamer:
    """Tests for services/dialog_namer.py."""

    def test_fallback_title_cleaning(self):
        text = "### **Помоги написать** скрипт на `Python` для парсинга"
        title = fallback_title_from_text(text)
        assert "###" not in title
        assert "**" not in title
        assert "`" not in title
        assert title[0].isupper()
        assert len(title) <= 35

    def test_fallback_title_empty_or_spaces(self):
        assert fallback_title_from_text("") == "Новый диалог"
        assert fallback_title_from_text("   \n\t  ") == "Новый диалог"
        assert fallback_title_from_text("### *** ```") == "Новый диалог"


class TestCalendarHelper:
    """Tests for services/calendar_helper.py."""

    def test_create_calendar_keyboard_ru(self):
        now = datetime.datetime.now()
        kb = create_calendar_keyboard(year=now.year, month=now.month, lang_code="ru")
        assert kb.inline_keyboard is not None
        assert len(kb.inline_keyboard) >= 5

        # Check navigation buttons
        has_date_button = False
        has_nav_button = False
        for row in kb.inline_keyboard:
            for btn in row:
                if btn.callback_data.startswith(CALLBACK_CALENDAR_DATE_PREFIX):
                    has_date_button = True
                if btn.callback_data.startswith(CALLBACK_CALENDAR_MONTH_PREFIX):
                    has_nav_button = True

        assert has_date_button is True
        assert has_nav_button is True

    def test_create_calendar_keyboard_en(self):
        now = datetime.datetime.now()
        kb = create_calendar_keyboard(year=now.year, month=now.month, lang_code="en")
        assert kb.inline_keyboard is not None

        # Verify English header
        header_text = kb.inline_keyboard[0][0].text
        # Header text should contain English month or year
        assert str(now.year) in header_text


@pytest.mark.asyncio
class TestGeminiServiceThinkingBudget:
    """Tests for thinking_budget configuration in GeminiService."""

    async def test_thinking_config_applied_to_gemini_2_5_models(self):
        from unittest.mock import AsyncMock, MagicMock, patch
        from services.gemini import GeminiService

        service = GeminiService(api_key="fake-test-key")

        mock_stream = AsyncMock()
        mock_chunk = MagicMock()
        mock_chunk.text = "Thought result"
        mock_stream.__aiter__.return_value = [mock_chunk]

        with patch.object(service.client.aio.models, "generate_content_stream", new=AsyncMock(return_value=mock_stream)) as mock_gen:
            # 1. Thinking model gemini-2.5-flash with budget 2048
            chunks = []
            async for c in service.generate_stream(
                model_id="gemini-2.5-flash",
                contents=[types.Content(role="user", parts=[types.Part.from_text(text="Hi")])],
                thinking_budget=2048,
            ):
                chunks.append(c)

            assert chunks == ["Thought result"]
            call_config = mock_gen.call_args.kwargs["config"]
            assert call_config.thinking_config is not None
            assert call_config.thinking_config.thinking_budget == 2048
            assert call_config.thinking_config.include_thoughts is True

            # 2. Non-thinking model (-lite) with budget provided -> thinking_config must be None
            chunks_lite = []
            async for c in service.generate_stream(
                model_id="gemini-2.5-flash-lite",
                contents=[types.Content(role="user", parts=[types.Part.from_text(text="Hi")])],
                thinking_budget=2048,
            ):
                chunks_lite.append(c)

            call_config_lite = mock_gen.call_args.kwargs["config"]
            assert call_config_lite.thinking_config is None


@pytest.mark.asyncio
class TestGeminiQuotaFallback:
    """Tests for 429 quota fallback tracking in GeminiService."""

    async def test_fallback_model_set_on_primary_fallback(self):
        from unittest.mock import AsyncMock, MagicMock, patch
        from google.genai.errors import APIError
        from services.gemini import GeminiService

        service = GeminiService(api_key="fake-test-key")
        assert service.fallback_model is None

        # First call fails with 429 RESOURCE_EXHAUSTED
        quota_err = APIError(429, "RESOURCE_EXHAUSTED: quota exceeded")

        # Fallback stream succeeds
        fallback_stream = AsyncMock()
        fallback_chunk = MagicMock()
        fallback_chunk.text = "Response from fallback"
        fallback_stream.__aiter__.return_value = [fallback_chunk]

        with patch.object(service.client.aio.models, "generate_content_stream", side_effect=[quota_err, fallback_stream]) as mock_gen:
            chunks = []
            async for c in service.generate_stream(
                model_id="gemini-3.5-flash",
                contents=[types.Content(role="user", parts=[types.Part.from_text(text="Calculate")])],
            ):
                chunks.append(c)

            assert chunks == ["Response from fallback"]
            assert service.fallback_model == "gemini-2.5-flash"
            assert mock_gen.await_count == 2
            # Second call was with gemini-2.5-flash
            assert mock_gen.call_args_list[1].kwargs["model"] == "gemini-2.5-flash"


